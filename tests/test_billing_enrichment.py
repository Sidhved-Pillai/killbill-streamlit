"""Regression coverage for missing rates and full addresses in statements."""

import io

import pandas as pd
from openpyxl import load_workbook

from billing_statement import build_billing_statement, build_billing_statement_workbook
from customer_master import apply_customer_master_lookup, build_customer_lookup
from freight_master import apply_freight_lookup, build_freight_lookup


def test_ocr_code_and_missing_loading_metadata_use_master_through_export():
    master = pd.DataFrame({
        "Unnamed: 3": ["MUMC019497"],
        "Name of the Distributor": ["RESHMA ENTERPRISES"],
        "Short Address": ["BHAYANDAR"],
        "Distance from Thane": [22],
        "Slab": ["21-30"],
        "Rate": [4508.64],
    })
    source = [{
        "Date": "05-Aug-26", "Invoice No.": "MUMCIN270048796",
        "Vehicle No.": "MH04HY7998", "From": "Thane",
        "Customer Code": "MUMCO19497", "Customer Name": "RESHMA ENTERPRISES",
        "To": "RESIMA ENTERPRISES COMPOUND NEAR DON BOSCO SCHOOL, PALGHAR 401105",
        "Loading Point": None, "Case": 1377, "Jar": 0,
    }]
    records = apply_customer_master_lookup(source, build_customer_lookup(master))
    records = apply_freight_lookup(records, build_freight_lookup(master))
    assert records[0]["Customer Code"] == "MUMC019497"
    assert records[0]["To"] == "BHAYANDAR"
    assert records[0]["Freight Charge"] == 4509
    assert source[0]["Customer Code"] == "MUMCO19497"
    statement = build_billing_statement(pd.DataFrame(records))
    sheet = load_workbook(io.BytesIO(build_billing_statement_workbook(statement))).active
    assert sheet["G4"].value == "BHAYANDAR"
    assert sheet["K4"].value == 4509
    assert sheet["M4"].value == 4509


def test_unknown_customer_is_not_matched_by_name_alone():
    source = [{"Customer Code": "MUMC999999", "Customer Name": "MARUTI ENTERPRISES", "To": "VASAI"}]
    lookup = {"MUMC018638": {"Customer Name": "MARUTI ENTERPRISES", "To": "AMBERNATH"}}
    result = apply_customer_master_lookup(source, lookup)[0]
    assert result["Lookup Status"] == "⚠ Not Found"
    assert result["To"] == "VASAI"


def test_ambiguous_ocr_code_is_not_guessed():
    lookup = {code: {"Customer Name": code, "To": "Short"} for code in ["MUMC001", "MUMCO01"]}
    result = apply_customer_master_lookup([{"Customer Code": "MUMCOO1", "To": "Full"}], lookup)[0]
    assert result["Lookup Status"] == "⚠ Not Found"
    assert result["To"] == "Full"


def test_missing_freight_stays_blank_and_totals_are_marked_partial():
    rows = [{
        "Date": "05-Aug-26", "Invoice No.": f"INV-{i}", "Vehicle No.": "TRUCK",
        "From": "Thane", "To": "Vasai", "Freight Charge": rate,
    } for i, rate in enumerate(["", 0, 4509])]
    statement = build_billing_statement(pd.DataFrame(rows))
    assert pd.isna(statement.iloc[0]["Freight"])
    assert statement.iloc[1]["Freight"] == 0
    sheet = load_workbook(io.BytesIO(build_billing_statement_workbook(statement))).active
    assert sheet["K4"].value is None
    assert sheet["M4"].value is None
    assert sheet["K5"].value == 0
    assert sheet["K24"].value == 4509
    assert sheet["A24"].value == "Partial Total (freight missing)"


def test_ambiguous_origin_does_not_assign_a_rate_and_metadata_takes_precedence():
    lookup = {"MUMC001": {"Thane": 4509, "Vasai": 3844}}
    records = apply_freight_lookup([
        {"Customer Code": "MUMC001", "From": "Thane / Vasai"},
        {"Customer Code": "MUMC001", "From": "Thane", "Loading Point": "Vasai"},
    ], lookup)
    assert records[0]["Freight Charge"] == ""
    assert records[1]["Freight Charge"] == 3844
