import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from billing_statement import build_billing_statement, build_billing_statement_workbook
from customer_master import apply_customer_master_lookup
from freight_master import apply_freight_lookup
from invoice_reference import apply_invoice_billing_reference


def invoice(**changes):
    record = {
        "Invoice No.": "MUMCIN270045533", "Date": "01-Aug-26",
        "From": "Bhiwandi", "Vehicle No.": "MH04LE8407",
        "Customer Name": "JAY GAVDEVI ENTERPRISES", "Customer Code": "UNKNOWN",
        "To": "HOUSE NO 1573 PETH VILLAGE ONE KHARGHAR, RAIGARH, Uran 410206",
        "Case": 0, "Jar": 400,
    }
    record.update(changes)
    return record


def test_reference_corrects_unmatched_invoice_through_excel_export():
    source = invoice()
    records = apply_customer_master_lookup([source], {})
    records = apply_freight_lookup(records, {})
    records = apply_invoice_billing_reference(records)
    result = records[0]
    assert result["To"] == "Kharghar"
    assert result["Freight Charge"] == 5073
    assert result["Lookup Status"] == "✅ Invoice reference"
    assert result["Jar"] == 400
    assert source["To"].startswith("HOUSE")
    summary = build_billing_statement(pd.DataFrame(records))
    sheet = load_workbook(io.BytesIO(build_billing_statement_workbook(summary))).active
    assert sheet["G4"].value == "Kharghar"
    assert sheet["K4"].value == 5073
    assert sheet["M24"].value == 5073


@pytest.mark.parametrize("changes", [
    {"Invoice No.": "MUMCIN999999999"},
    {"Date": "01-Aug-27"},
    {"Date": ""},
    {"From": "Thane"},
    {"From": ""},
])
def test_same_customer_does_not_receive_unrelated_invoice_rate(changes):
    source = invoice(**changes)
    assert apply_invoice_billing_reference([source]) == [source]


def test_numeric_invoice_gets_reference_rate_without_ton_prefix():
    result = apply_invoice_billing_reference([invoice(**{
        "Invoice No.": "292264017722", "Date": "12-Aug-26",
    })])[0]
    assert result["Freight Charge"] == 4327
    assert result["To"] == "Bhiwandi"


def test_reference_fills_missing_rate_and_handles_normalized_keys():
    result = apply_invoice_billing_reference([invoice(**{
        "Invoice No.": " mumcin270045533 ", "Date": "2026-08-01",
        "From": "Bhiwandi Depot", "To": "Old location", "Freight Charge": "",
    })])[0]
    assert result["Freight Charge"] == 5073
    assert result["To"] == "Kharghar"


def test_updated_master_values_take_precedence_over_positive_photo_rates():
    source = invoice(**{
        "Freight Charge": 3844, "To": "Updated master location",
        "Lookup Status": "✅ Matched",
    })
    assert apply_invoice_billing_reference([source]) == [source]


def test_photo_fills_missing_rate_without_replacing_master_location():
    source = invoice(**{
        "Freight Charge": "", "To": "Updated master location",
        "Lookup Status": "✅ Matched",
    })
    result = apply_invoice_billing_reference([source])[0]
    assert result["Freight Charge"] == 5073
    assert result["To"] == "Updated master location"


def test_reshma_uses_confirmed_location():
    result = apply_invoice_billing_reference([invoice(**{
        "Invoice No.": "MUMCIN270048796", "Date": "05-Aug-26",
        "From": "Thane", "To": "Extracted address",
    })])[0]
    assert result["Freight Charge"] == 4509
    assert result["To"] == "Bhayander E"
    assert result["Lookup Status"] == "✅ Invoice reference"


@pytest.mark.parametrize("number,date,location", [
    ("MUMCIN270053426", "23-Aug-26", "Kopar Khairane"),
    ("MUMCIN270049081", "07-Aug-26", "SHREE SWAMI COLDDRINKS"),
])
@pytest.mark.parametrize("previous_rate", ["", 5073])
def test_confirmed_zero_rates_remain_numeric_zero_in_export(number, date, location, previous_rate):
    source = invoice(**{
        "Invoice No.": number, "Date": date, "Freight Charge": previous_rate,
        "Lookup Status": "✅ Matched",
    })
    records = apply_invoice_billing_reference([source])
    assert records[0]["Freight Charge"] == 0
    assert records[0]["To"] == location
    assert records[0]["Lookup Status"] == "✅ Invoice reference"
    summary = build_billing_statement(pd.DataFrame(records))
    assert not summary["Freight"].isna().any()
    sheet = load_workbook(io.BytesIO(build_billing_statement_workbook(summary))).active
    assert sheet["G4"].value == location
    assert sheet["K4"].value == 0
    assert sheet["M4"].value == 0
    assert sheet["M24"].value == 0
    assert sheet["A24"].value == "Grand Total =======>"
